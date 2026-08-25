import fitz
import pytest

from ragkb.models import TableData
from ragkb.parsers.base import parse_document, table_to_markdown
from ragkb.parsers.pdf_parser import PdfParser


@pytest.fixture
def sample_pdf(tmp_path):
    path = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    # 注意：insert_text 默认用 Helvetica，无法编码中文（提取出来是乱码），
    # 必须显式指定内置中文字体 china-s（简体）。
    page.insert_text((72, 72), "产品规格说明。", fontname="china-s")
    page.insert_text((72, 120), "型号 A-100 单价 99 元。", fontname="china-s")
    rect = fitz.Rect(72, 160, 200, 200)
    # 1x1 RGBA PNG（有效字节流；计划原文的 hex 截断了 IEND CRC，
    # 会被 PyMuPDF 以 "premature end of data in png image" 拒绝）
    page.insert_image(rect, stream=bytes.fromhex("89504E470D0A1A0A0000000D49484452000000010000000108060000001F15C4890000000B49444154789C6360000200000500017A5EAB3F0000000049454E44AE426082"))
    doc.save(path)
    doc.close()
    return path


def _draw_table(page, x0, y0, n_cols, n_rows, cells):
    """在页面上画 n_cols 列 x n_rows 行（含表头）的表格线网格并填入文字。

    fitz.draw_line 画出的矢量线会被 pdfplumber 的 lines 策略识别为表格
    边界（本环境实测 pdfplumber 0.11.10 + PyMuPDF 1.28.2 可稳定检出）。
    """
    xs = [x0 + i * 80 for i in range(n_cols + 1)]
    ys = [y0 + i * 30 for i in range(n_rows + 1)]
    for x in xs:
        page.draw_line(fitz.Point(x, ys[0]), fitz.Point(x, ys[-1]),
                       color=(0, 0, 0), width=1)
    for y in ys:
        page.draw_line(fitz.Point(xs[0], y), fitz.Point(xs[-1], y),
                       color=(0, 0, 0), width=1)
    for r in range(n_rows):
        for c in range(n_cols):
            page.insert_text(fitz.Point(xs[c] + 5, ys[r] + 22), cells[r][c],
                             fontname="china-s", fontsize=10)


def test_pdf_parser_extracts_text_and_image(sample_pdf):
    result = PdfParser().parse(str(sample_pdf), doc_id="d1", source="sample.pdf")
    assert "产品规格说明" in result.text
    assert len(result.images) >= 1
    assert result.images[0].data[:8] == b"\x89PNG\r\n\x1a\n"


def test_pdf_parser_extracts_tables_with_unique_ids(tmp_path):
    # 同一页两张尺寸完全相同的表格（2 列 x 3 行，含表头）——
    # 修复前 table_id 只含 页号+维度，两张会撞成同一个 id；
    # 现在按页内序号区分，验证 table_id 唯一且内容按序落位。
    # 第二张表末格留空：走解析循环的 None/空串 -> "N/A" 转换。
    path = tmp_path / "tables.pdf"
    doc = fitz.open()
    page = doc.new_page()
    _draw_table(page, 72, 100, 2, 3,
                [("型号", "A-100"), ("单价", "99"), ("库存", "10")])
    _draw_table(page, 72, 250, 2, 3,
                [("型号", "B-200"), ("单价", "199"), ("库存", "")])
    doc.save(path)
    doc.close()

    result = PdfParser().parse(str(path), doc_id="d1", source="tables.pdf")

    assert len(result.tables) == 2
    ids = [t.table_id for t in result.tables]
    assert len(set(ids)) == len(ids)  # table_id 全局唯一
    # 表头取第一行；行数据为其余行
    assert result.tables[0].headers == ["型号", "A-100"]
    assert result.tables[0].rows == [["单价", "99"], ["库存", "10"]]
    assert result.tables[1].headers == ["型号", "B-200"]
    assert result.tables[1].rows == [["单价", "199"], ["库存", "N/A"]]
    # 页内序号参与 id 与名称，两张同维表格不再冲突
    assert ids[0] == "tables.pdf-1-1-2x2"
    assert ids[1] == "tables.pdf-1-2-2x2"
    assert result.tables[0].name == "第1页第1张表格"
    assert result.tables[1].name == "第1页第2张表格"


def test_make_table_id_unique_by_tbl_idx():
    # 单元级锁定 _make_table 契约（静态方法，N/A 转换由解析循环负责）：
    # 同页同维度的两张表仅靠 tbl_idx 区分 id，互不冲突。
    t1 = PdfParser._make_table(["型号", "单价"],
                               [["A", None], [None, "99"]], "u.pdf", 1, 1)
    t2 = PdfParser._make_table(["型号", "单价"],
                               [["B", None], [None, "199"]], "u.pdf", 1, 2)
    assert t1.table_id == "u.pdf-1-1-2x2"
    assert t2.table_id == "u.pdf-1-2-2x2"
    assert t1.table_id != t2.table_id
    assert t1.headers == ["型号", "单价"]
    assert t1.source == "u.pdf:1"
    assert isinstance(t1, TableData)


def test_parse_document_routes_pdf_and_passes_metadata(sample_pdf):
    result = parse_document(str(sample_pdf), doc_id="d1", source="sample.pdf",
                            department="质量部", version="v2.1",
                            effective_date="2025-06-01")
    assert result.doc_type == "pdf"
    assert result.doc_id == "d1"
    assert result.source == "sample.pdf"
    assert result.department == "质量部"
    assert result.version == "v2.1"
    assert result.effective_date == "2025-06-01"


def test_parse_document_rejects_unsupported_extension(tmp_path):
    path = tmp_path / "note.txt"
    path.write_text("hello", encoding="utf-8")
    with pytest.raises(ValueError, match="不支持的文档类型"):
        parse_document(str(path), doc_id="d1", source="note.txt")


# ---- Task 4: Excel 与邮件解析器 ----

from email.message import EmailMessage

from openpyxl import Workbook

from ragkb.parsers.email_parser import EmailParser
from ragkb.parsers.excel_parser import ExcelParser


def _make_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "报价"
    ws.append(["型号", "单价"])
    ws.append(["A-100", "99"])
    wb.save(path)


def _make_eml(path):
    # 用 stdlib EmailMessage 生成真实格式的 .eml：Content-Type 带 charset，
    # Subject 走 RFC 2047 编码——裸写 UTF-8 文本会被 stdlib 按 ASCII 解码
    # 成 U+FFFD（无 charset 头时 get_content() 以 ascii+replace 解码）。
    msg = EmailMessage()
    msg["From"] = "a@x.com"
    msg["To"] = "b@x.com"
    msg["Subject"] = "报价更新"
    msg.set_content("新单价见附件。")
    path.write_bytes(msg.as_bytes())


def test_excel_parser_keeps_row_column_structure(tmp_path):
    path = tmp_path / "t.xlsx"
    _make_xlsx(path)
    result = ExcelParser().parse(str(path), doc_id="d2", source="t.xlsx")
    assert result.doc_type == "excel"
    assert result.tables[0].headers == ["型号", "单价"]
    assert result.tables[0].rows == [["A-100", "99"]]
    assert result.tables[0].source == "t.xlsx:报价"


def test_table_to_markdown_serializes_headers_and_rows():
    # 表格 A 路：TableData → Markdown 文本，供语义检索（与正文分块独立）
    md = table_to_markdown(TableData(
        table_id="t.xlsx-报价", name="报价", headers=["型号", "单价"],
        rows=[["A-100", "99"]], source="t.xlsx:报价"))
    lines = md.splitlines()
    assert lines[0] == "## 表格：报价"
    assert "来源：t.xlsx:报价" in md
    assert "| 型号 | 单价 |" in md
    assert "| --- | --- |" in md
    assert "| A-100 | 99 |" in md


def test_email_parser_extracts_body(tmp_path):
    path = tmp_path / "m.eml"
    _make_eml(path)
    result = EmailParser().parse(str(path), doc_id="d3", source="m.eml")
    assert "新单价" in result.text


def test_parse_document_routes_excel_and_email(tmp_path):
    # 本任务补齐 base.parse_document 的 xlsx/eml 惰性导入分支：
    # 分发结果 + 元数据透传。
    xlsx = tmp_path / "t.xlsx"
    _make_xlsx(xlsx)
    eml = tmp_path / "m.eml"
    _make_eml(eml)
    excel = parse_document(str(xlsx), doc_id="d2", source="t.xlsx",
                           department="质量部", version="v1",
                           effective_date="2025-06-01")
    email_doc = parse_document(str(eml), doc_id="d3", source="m.eml",
                               department="市场部", version="v2",
                               effective_date="2025-07-01")
    assert excel.doc_type == "excel"
    assert excel.tables[0].headers == ["型号", "单价"]
    assert excel.department == "质量部"
    assert excel.version == "v1"
    assert excel.effective_date == "2025-06-01"
    assert email_doc.doc_type == "email"
    assert "新单价" in email_doc.text
    assert email_doc.department == "市场部"
    assert email_doc.version == "v2"
    assert email_doc.effective_date == "2025-07-01"


def test_parse_document_rejects_xls_extension(tmp_path):
    # 旧 .xls 格式 openpyxl 无法解析，已从分发映射移除——
    # 应落入不支持的扩展名分支抛 ValueError，而不是硬失败。
    path = tmp_path / "old.xls"
    path.write_bytes(b"dummy xls bytes")
    with pytest.raises(ValueError, match="不支持的文档类型"):
        parse_document(str(path), doc_id="d1", source="old.xls")


def test_excel_parser_without_headers(tmp_path):
    # has_headers=False：第一行不再被当作表头，
    # 所有行都进 rows，headers 为空列表。
    path = tmp_path / "t_noheader.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["A-100", "99"])
    ws.append(["B-200", "199"])
    wb.save(path)
    result = ExcelParser().parse(str(path), doc_id="d4",
                                 source="t_noheader.xlsx", has_headers=False)
    assert result.tables[0].headers == []
    assert result.tables[0].rows == [["A-100", "99"], ["B-200", "199"]]


def test_email_parser_html_only_body(tmp_path):
    # 纯 HTML 邮件（无 text/plain 部分）：回退提取 text/html 并去标签。
    path = tmp_path / "m_html.eml"
    msg = EmailMessage()
    msg["From"] = "a@x.com"
    msg["To"] = "b@x.com"
    msg["Subject"] = "公告"
    msg.add_alternative(
        "<html><body><p>重要通知已发布</p></body></html>", subtype="html")
    path.write_bytes(msg.as_bytes())
    result = EmailParser().parse(str(path), doc_id="d5", source="m_html.eml")
    assert "重要通知已发布" in result.text


def test_email_parser_joins_multiple_plain_parts(tmp_path):
    # 多个 text/plain part 以换行分隔拼接，而非直接 += 粘连。
    from email.mime.multipart import MIMEMultipart
    from email.mime.text import MIMEText

    path = tmp_path / "m_multi.eml"
    msg = MIMEMultipart()
    msg["From"] = "a@x.com"
    msg["To"] = "b@x.com"
    msg["Subject"] = "多段正文"
    msg.attach(MIMEText("第一段正文", "plain", "utf-8"))
    msg.attach(MIMEText("第二段正文", "plain", "utf-8"))
    path.write_bytes(msg.as_bytes())
    result = EmailParser().parse(str(path), doc_id="d6", source="m_multi.eml")
    assert "第一段正文" in result.text
    assert "第二段正文" in result.text


def test_email_parser_omits_empty_subject_line(tmp_path):
    # Subject 为空时不输出 "主题：" 行。
    path = tmp_path / "m_nosubject.eml"
    msg = EmailMessage()
    msg["From"] = "a@x.com"
    msg["To"] = "b@x.com"
    msg.set_content("正文内容")
    path.write_bytes(msg.as_bytes())
    result = EmailParser().parse(str(path), doc_id="d7", source="m_nosubject.eml")
    assert "主题：" not in result.text
    assert "正文内容" in result.text
